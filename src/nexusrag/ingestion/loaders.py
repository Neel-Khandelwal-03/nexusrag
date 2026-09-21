"""Parse files and web pages into structured :class:`~nexusrag.models.Document` objects.

Loaders only extract *structure* (headings, paragraphs, lists, tables and code, with
page numbers where the format has them); every chunking decision lives in
``chunker.py``. Everything here is synchronous and pure except :func:`fetch_url`,
which does network I/O.

Format notes:

* **PDF**: text and font sizes come from PyMuPDF, tables from pdfplumber (rendered as
  Markdown). Text inside a detected table is dropped from the text stream so it isn't
  indexed twice. Headings come from the PDF outline when there is one, otherwise from
  a font-size heuristic. Repeated running headers/footers and bare page numbers are
  removed. Scanned (image-only) PDFs have no text layer and are rejected (no OCR).
* **DOCX**: headings from paragraph styles ("Heading N"), tables in document order.
* **Markdown**: a small hand-written parser for ATX/setext headings, fenced code,
  pipe tables and lists. Front-matter ``title:`` is honoured.
* **TXT**: paragraphs split on blank lines, with conservative heading detection
  (ALL-CAPS lines and numbered headings like ``2.1 Data``).
* **URL**: fetched with SSRF protection, main content extracted by trafilatura as
  Markdown, then parsed like a Markdown file.
"""

from __future__ import annotations

import asyncio
import ipaddress
import math
import re
import socket
from collections import Counter
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlparse

import httpx

from nexusrag.ingestion.metadata import (
    clean_text,
    file_sha256,
    infer_title,
    make_doc_id,
    normalize_source_key,
    sha256_text,
    title_from_filename,
)
from nexusrag.log import get_logger
from nexusrag.models import Document, DocumentElement, ElementKind, SourceType

log = get_logger(__name__)

SUPPORTED_EXTENSIONS: dict[str, SourceType] = {
    ".pdf": SourceType.PDF,
    ".docx": SourceType.DOCX,
    ".md": SourceType.MARKDOWN,
    ".markdown": SourceType.MARKDOWN,
    ".txt": SourceType.TEXT,
}

USER_AGENT = "NexusRAG/0.1 (+https://github.com/Neel-Khandelwal-03/nexusrag)"
MAX_REDIRECTS = 5


class LoaderError(ValueError):
    """A source couldn't be loaded. The message is safe to show to users."""


class UnsupportedFileError(LoaderError):
    """The file type isn't supported."""


def detect_source_type(filename: str) -> SourceType:
    suffix = Path(filename).suffix.lower()
    try:
        return SUPPORTED_EXTENSIONS[suffix]
    except KeyError:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFileError(
            f"Unsupported file type {suffix or '(none)'!r} for {filename}. Supported: {supported}."
        ) from None


def _heading(text: str, level: int, page: int | None = None) -> DocumentElement:
    return DocumentElement(
        kind=ElementKind.HEADING, text=text, level=max(1, min(level, 6)), page=page
    )


def _element(kind: ElementKind, text: str, page: int | None = None) -> DocumentElement:
    return DocumentElement(kind=kind, text=text, page=page)


# --------------------------------------------------------------------------- tables


def _cell(value: Any) -> str:
    return " ".join(str(value or "").split()).replace("|", "\\|")


def table_to_markdown(rows: Sequence[Sequence[Any]], *, min_rows: int = 2) -> str:
    """Render rows (first row = header) as a GitHub-flavoured Markdown table.

    Returns "" for anything that isn't really a table (fewer than ``min_rows`` rows or
    < 2 columns), so callers can treat that content as ordinary text instead.
    """
    cleaned = [[_cell(c) for c in row] for row in rows if row and any(_cell(c) for c in row)]
    if len(cleaned) < min_rows or not cleaned:
        return ""
    width = max(len(row) for row in cleaned)
    if width < 2:
        return ""
    padded = [row + [""] * (width - len(row)) for row in cleaned]
    header, *body = padded
    lines = ["| " + " | ".join(header) + " |", "|" + "|".join(["---"] * width) + "|"]
    lines += ["| " + " | ".join(row) + " |" for row in body]
    return "\n".join(lines)


# --------------------------------------------------------------------------- markdown

_ATX_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_FENCE_RE = re.compile(r"^(`{3,}|~{3,})")
_LIST_RE = re.compile(r"^\s*(?:[-*+•▪◦‣]|\d{1,3}[.)])\s+")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)*\|?\s*$")
_SETEXT_RE = re.compile(r"^(?:=+|-+)\s*$")
_HR_RE = re.compile(r"^(?:(?:\*\s*){3,}|(?:-\s*){3,}|(?:_\s*){3,})$")
_FRONT_MATTER_TITLE_RE = re.compile(r"^title:\s*[\"']?(.+?)[\"']?\s*$", re.MULTILINE)


def _strip_front_matter(text: str) -> tuple[str, str | None]:
    if not text.startswith("---\n"):
        return text, None
    end = text.find("\n---", 4)
    if end == -1:
        return text, None
    match = _FRONT_MATTER_TITLE_RE.search(text[4:end])
    return text[end + 4 :].lstrip("-\n"), match.group(1).strip() if match else None


def parse_markdown(text: str) -> tuple[list[DocumentElement], str | None]:
    """Parse Markdown into elements. Returns ``(elements, front_matter_title)``."""
    text, title = _strip_front_matter(clean_text(text, preserve_whitespace=True))
    lines = text.split("\n")
    elements: list[DocumentElement] = []
    buffer: list[str] = []
    buffer_kind: ElementKind | None = None

    def flush() -> None:
        nonlocal buffer_kind
        if buffer and buffer_kind is not None:
            is_list = buffer_kind == ElementKind.LIST
            # Lists keep their indentation (nested items); paragraphs are reflowed.
            raw = "\n".join(buffer) if is_list else " ".join(buffer)
            joined = clean_text(raw, preserve_whitespace=is_list)
            if joined:
                elements.append(_element(buffer_kind, joined))
        buffer.clear()
        buffer_kind = None

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        fence = _FENCE_RE.match(stripped)
        if fence:
            flush()
            marker = fence.group(1)
            code = [line]
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(marker):
                code.append(lines[i])
                i += 1
            code.append(marker)
            i += 1
            elements.append(_element(ElementKind.CODE, "\n".join(code)))
            continue

        atx = _ATX_RE.match(stripped)
        if atx:
            flush()
            elements.append(_heading(clean_text(atx.group(2)), len(atx.group(1))))
            i += 1
            continue

        # Setext heading: a paragraph line underlined with === (h1) or --- (h2).
        if stripped and buffer_kind == ElementKind.PARAGRAPH and _SETEXT_RE.match(stripped):
            heading_text = clean_text(" ".join(buffer))
            buffer.clear()
            buffer_kind = None
            elements.append(_heading(heading_text, 1 if stripped[0] == "=" else 2))
            i += 1
            continue

        if not stripped:
            flush()
            i += 1
            continue

        if _HR_RE.match(stripped):
            flush()
            i += 1
            continue

        # Pipe table: a row followed by a separator row.
        if "|" in stripped and i + 1 < len(lines) and _TABLE_SEP_RE.match(lines[i + 1]):
            flush()
            rows = [stripped, lines[i + 1].strip()]
            i += 2
            while i < len(lines) and lines[i].strip() and "|" in lines[i]:
                rows.append(lines[i].strip())
                i += 1
            elements.append(_element(ElementKind.TABLE, "\n".join(rows)))
            continue

        if _LIST_RE.match(line):
            if buffer_kind != ElementKind.LIST:
                flush()
                buffer_kind = ElementKind.LIST
            buffer.append(line.rstrip())
            i += 1
            continue

        if buffer_kind == ElementKind.LIST and line[:1] in (" ", "\t"):
            buffer.append(line.rstrip())  # continuation of a list item
            i += 1
            continue

        if buffer_kind != ElementKind.PARAGRAPH:
            flush()
            buffer_kind = ElementKind.PARAGRAPH
        buffer.append(stripped.lstrip("> ").strip() if stripped.startswith(">") else stripped)
        i += 1

    flush()
    return elements, title


# --------------------------------------------------------------------------- plain text

_NUMBERED_HEADING_RE = re.compile(r"^(\d{1,2}(?:\.\d{1,2}){0,4})\.?\s+([A-Z].{0,78})$")


def _plain_heading_level(block: str) -> int | None:
    """Heading level for a one-line block that looks like a heading, else None."""
    if "\n" in block or len(block) > 80 or block.endswith((".", ",", ";", "?", "!")):
        return None
    numbered = _NUMBERED_HEADING_RE.match(block)
    if numbered:
        return numbered.group(1).count(".") + 1
    letters = [ch for ch in block if ch.isalpha()]
    if len(letters) >= 3 and all(ch.isupper() for ch in letters):
        return 1
    return None


def parse_plain_text(text: str) -> list[DocumentElement]:
    """Split plain text into paragraphs, lists and (conservatively detected) headings."""
    elements: list[DocumentElement] = []
    for raw in re.split(r"\n\s*\n", clean_text(text)):
        block = raw.strip()
        if not block:
            continue
        level = _plain_heading_level(block)
        if level is not None:
            elements.append(_heading(block, level))
            continue
        lines = block.split("\n")
        if all(_LIST_RE.match(line) for line in lines):
            elements.append(_element(ElementKind.LIST, block))
        else:
            elements.append(_element(ElementKind.PARAGRAPH, " ".join(lines)))
    return elements


def read_text_file(path: Path) -> str:
    """Read a text file as UTF-8 (BOM tolerated), falling back to cp1252."""
    data = path.read_bytes()
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


# --------------------------------------------------------------------------- docx


def load_docx_elements(path: Path) -> tuple[list[DocumentElement], str | None]:
    """Extract headings, paragraphs, lists and tables from a .docx in document order."""
    import docx  # lazy: heavy import only when needed
    from docx.table import Table

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise LoaderError(f"Could not open {path.name} as a Word document.") from exc

    title = (document.core_properties.title or "").strip() or None
    elements: list[DocumentElement] = []
    list_items: list[str] = []

    def flush_list() -> None:
        if list_items:
            elements.append(_element(ElementKind.LIST, "\n".join(list_items)))
            list_items.clear()

    for block in document.iter_inner_content():
        if isinstance(block, Table):
            flush_list()
            markdown = table_to_markdown([[cell.text for cell in row.cells] for row in block.rows])
            if markdown:
                elements.append(_element(ElementKind.TABLE, markdown))
            continue
        text = clean_text(block.text)
        if not text:
            continue
        style = (block.style.name if block.style is not None else "").lower()
        if style == "title":
            # The document title is metadata (and the embedding prefix), not a section.
            flush_list()
            title = title or text
        elif style.startswith("heading"):
            flush_list()
            digits = re.findall(r"\d", style)
            elements.append(_heading(text, int(digits[0]) if digits else 1))
        elif "list" in style:
            list_items.append(f"- {text}")
        else:
            flush_list()
            elements.append(_element(ElementKind.PARAGRAPH, text))
    flush_list()
    return elements, title


# --------------------------------------------------------------------------- pdf

_PAGE_NUMBER_RE = re.compile(r"^(?:page\s*)?\d{1,4}(?:\s*(?:/|of)\s*\d{1,4})?$", re.IGNORECASE)
_PAGE_LABEL_RE = re.compile(r"\bpage\s+\d{1,4}\b", re.IGNORECASE)
_BULLET_RE = re.compile(r"^\s*(?:[•▪◦‣●○■□–-]|\d{1,3}[.)])\s+")
_BULLET_CHARS = "•▪◦‣●○■□–"


BBox = tuple[float, float, float, float]
PageTables = dict[int, list[tuple[BBox, str]]]  # page -> [(bbox, markdown)]


@dataclass
class _PdfLine:
    page: int
    block: int
    top: float
    rel_top: float  # position on the page, 0 = top edge, 1 = bottom edge
    text: str
    size: float
    bold: bool


def _normalize_for_match(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _extract_pdf_tables(
    path: Path,
) -> PageTables:
    """Tables per page as ``(bbox, markdown)``. Failures degrade to "no tables"."""
    import pdfplumber

    tables: PageTables = {}
    try:
        with pdfplumber.open(str(path)) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                for table in page.find_tables():
                    # Keep header-only fragments: a table broken across pages leaves its
                    # header at the bottom of one page, merged back later.
                    markdown = table_to_markdown(table.extract(), min_rows=1)
                    if markdown:
                        tables.setdefault(page_no, []).append((tuple(table.bbox), markdown))  # type: ignore[arg-type]
    except Exception as exc:  # pdfplumber is best-effort; text extraction still works
        log.warning("loader.pdf_tables_failed", filename=path.name, error=type(exc).__name__)
    return tables


def _pdf_lines(doc: Any, tables: PageTables) -> list[_PdfLine]:
    import pymupdf

    # Ligatures (ﬁ, ﬂ) would break keyword search, and images aren't needed.
    flags = (
        pymupdf.TEXTFLAGS_DICT & ~pymupdf.TEXT_PRESERVE_LIGATURES & ~pymupdf.TEXT_PRESERVE_IMAGES
    )
    lines: list[_PdfLine] = []
    for page_no, page in enumerate(doc, start=1):
        height = float(page.rect.height) or 1.0
        boxes = [bbox for bbox, _ in tables.get(page_no, [])]
        data = page.get_text("dict", sort=True, flags=flags)
        for block_no, block in enumerate(data.get("blocks", [])):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
                if not spans:
                    continue
                x0, y0, x1, y1 = line["bbox"]
                cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
                if any(b[0] <= cx <= b[2] and b[1] <= cy <= b[3] for b in boxes):
                    continue  # rendered from pdfplumber's table instead
                text = clean_text("".join(s["text"] for s in spans))
                if not text:
                    continue
                lines.append(
                    _PdfLine(
                        page=page_no,
                        block=block_no,
                        top=float(y0),
                        rel_top=float(y0) / height,
                        text=text,
                        size=round(max(float(s["size"]) for s in spans), 1),
                        bold=all(
                            (int(s.get("flags", 0)) & 16)
                            or "bold" in str(s.get("font", "")).lower()
                            for s in spans
                        ),
                    )
                )
    return lines


def _in_margin(line: _PdfLine) -> bool:
    return line.rel_top < 0.1 or line.rel_top > 0.9


def _drop_page_furniture(lines: list[_PdfLine], page_count: int) -> list[_PdfLine]:
    """Remove page numbers and running headers/footers.

    Dropped: bare page numbers anywhere, margin lines containing "Page N", and margin
    lines that repeat (digits ignored) on at least half the pages of a multi-page PDF.
    """

    def key(line: _PdfLine) -> str:
        return re.sub(r"\d+", "#", _normalize_for_match(line.text))

    furniture: set[str] = set()
    if page_count >= 2:
        pages_per_key: dict[str, set[int]] = {}
        for ln in lines:
            if _in_margin(ln):
                pages_per_key.setdefault(key(ln), set()).add(ln.page)
        threshold = max(2, math.ceil(page_count / 2))
        furniture = {k for k, pages in pages_per_key.items() if len(pages) >= threshold}
    return [
        ln
        for ln in lines
        if not _PAGE_NUMBER_RE.match(ln.text)
        and not (_in_margin(ln) and (key(ln) in furniture or _PAGE_LABEL_RE.search(ln.text)))
    ]


def _body_font_size(lines: Sequence[_PdfLine]) -> float:
    weights: Counter[float] = Counter()
    for ln in lines:
        weights[ln.size] += len(ln.text)
    return weights.most_common(1)[0][0] if weights else 0.0


def _heading_levels(lines: Sequence[_PdfLine], toc: Sequence[Sequence[Any]]) -> dict[int, int]:
    """Map line index -> heading level, from the outline if present else font sizes."""
    levels: dict[int, int] = {}
    if toc:
        wanted: dict[tuple[int, str], int] = {}
        for entry in toc:
            level, title, page = int(entry[0]), str(entry[1]), int(entry[2])
            for p in (page - 1, page, page + 1):
                wanted.setdefault((p, _normalize_for_match(title)), level)
        for i, ln in enumerate(lines):
            found = wanted.get((ln.page, _normalize_for_match(ln.text)))
            if found is not None:
                levels[i] = found
        if levels:
            return levels

    body = _body_font_size(lines)
    if body <= 0:
        return levels

    def is_candidate(ln: _PdfLine) -> bool:
        text = ln.text
        if (
            len(text) > 120
            or not any(ch.isalpha() for ch in text)
            or text.endswith((".", ",", ";"))
        ):
            return False
        larger = ln.size >= body * 1.15
        if _BULLET_RE.match(text) and not larger:
            return (
                False  # a bold list item, not a heading ("1. Introduction" in a big font is fine)
            )
        return larger or (ln.bold and ln.size >= body and len(text) <= 80)

    candidates = [i for i, ln in enumerate(lines) if is_candidate(lines[i])]
    larger_sizes = sorted(
        {lines[i].size for i in candidates if lines[i].size >= body * 1.15}, reverse=True
    )
    for i in candidates:
        size = lines[i].size
        levels[i] = larger_sizes.index(size) + 1 if size in larger_sizes else len(larger_sizes) + 1
    return levels


def _join_lines(parts: Sequence[str]) -> str:
    """Join wrapped lines into one string.

    A line ending in "-" is joined *without* a space but the hyphen is kept: we can't
    tell a word hyphenated by the typesetter ("inspec-tion") from a real compound broken
    at its hyphen ("wind-turbine"), and dropping a real hyphen corrupts the word, while
    keeping a soft one only splits it for keyword search.
    """
    out = ""
    for part in parts:
        if out.endswith("-") and part[:1].islower():
            out += part
        else:
            out = f"{out} {part}" if out else part
    return out


def _table_width(markdown: str) -> int:
    return markdown.split("\n", 1)[0].count("|") - 1


def _merge_page_continuations(elements: list[DocumentElement]) -> list[DocumentElement]:
    """Re-join paragraphs and tables that a page break split in two.

    * A paragraph at the end of a page that doesn't end a sentence continues in the
      first paragraph of the next page.
    * A table at the top of a page right after a table with the same number of columns
      at the end of the previous page is its continuation. A repeated header row is
      dropped; otherwise the continuation's first row is data.
    """
    merged: list[DocumentElement] = []
    for element in elements:
        prev = merged[-1] if merged else None
        if (
            prev is None
            or prev.page is None
            or element.page != prev.page + 1
            or prev.kind != element.kind
        ):
            merged.append(element)
            continue
        unfinished = not prev.text.endswith((".", "!", "?", ":", ";"))
        if element.kind == ElementKind.PARAGRAPH and unfinished:
            text = _join_lines([prev.text, element.text])
            merged[-1] = _element(ElementKind.PARAGRAPH, text, prev.page)
            continue
        if element.kind == ElementKind.TABLE and _table_width(prev.text) == _table_width(
            element.text
        ):
            prev_lines, lines = prev.text.split("\n"), element.text.split("\n")
            rows = lines[2:] if lines[0] == prev_lines[0] else [lines[0], *lines[2:]]
            table = "\n".join([*prev_lines, *rows])
            merged[-1] = _element(ElementKind.TABLE, table, prev.page)
            continue
        merged.append(element)
    return merged


def _pdf_elements(
    lines: Sequence[_PdfLine],
    levels: dict[int, int],
    tables: PageTables,
) -> list[DocumentElement]:
    elements: list[DocumentElement] = []
    paragraph: list[str] = []
    items: list[str] = []
    # "start" is the page where the paragraph/list being built began.
    state: dict[str, Any] = {"start": None, "block": None}

    def flush() -> None:
        if paragraph:
            elements.append(_element(ElementKind.PARAGRAPH, _join_lines(paragraph), state["start"]))
            paragraph.clear()
        if items:
            elements.append(_element(ElementKind.LIST, "\n".join(items), state["start"]))
            items.clear()

    pending_tables = {page: sorted(ts, key=lambda t: t[0][1]) for page, ts in tables.items()}

    def emit_tables(page: int, before: float | None) -> None:
        queue = pending_tables.get(page, [])
        while queue and (before is None or queue[0][0][1] <= before):
            flush()
            elements.append(_element(ElementKind.TABLE, queue.pop(0)[1], page))

    last_page = 0
    for i, ln in enumerate(lines):
        if ln.page != last_page:
            for page in range(last_page, ln.page):
                emit_tables(page, None)  # tables left on earlier pages
            last_page = ln.page
        emit_tables(ln.page, ln.top)

        level = levels.get(i)
        if level is not None:
            flush()
            previous = elements[-1] if elements else None
            # Merge a heading wrapped over two lines of the same block.
            if (
                previous is not None
                and previous.kind == ElementKind.HEADING
                and previous.level == level
                and state["block"] == (ln.page, ln.block)
            ):
                elements[-1] = _heading(f"{previous.text} {ln.text}", level, previous.page)
            else:
                elements.append(_heading(ln.text, level, ln.page))
            state.update(block=(ln.page, ln.block))
            continue

        new_block = state["block"] != (ln.page, ln.block)
        if _BULLET_RE.match(ln.text):
            # Consecutive bullets stay one list even when each sits in its own PDF block.
            if paragraph:
                flush()
            if not items:
                state["start"] = ln.page
            text = ln.text.lstrip(_BULLET_CHARS + " ").strip()
            items.append(f"- {text}" if ln.text.lstrip()[:1] in _BULLET_CHARS + "-" else ln.text)
        elif items and not new_block:
            items[-1] = _join_lines([items[-1], ln.text])  # wrapped list item
        else:
            if new_block:
                flush()
            if not paragraph:
                state["start"] = ln.page
            paragraph.append(ln.text)
        state.update(block=(ln.page, ln.block))

    flush()
    for page in sorted(pending_tables):
        emit_tables(page, None)
    return elements


def load_pdf_elements(path: Path) -> tuple[list[DocumentElement], str | None]:
    """Extract structured elements (with 1-based page numbers) from a PDF."""
    import pymupdf

    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise LoaderError(f"Could not open {path.name} as a PDF.") from exc
    with doc:
        if doc.needs_pass:
            raise LoaderError(f"{path.name} is password-protected.")
        title = ((doc.metadata or {}).get("title") or "").strip() or None
        tables = _extract_pdf_tables(path)
        lines = _drop_page_furniture(_pdf_lines(doc, tables), doc.page_count)
        levels = _heading_levels(lines, doc.get_toc(simple=True))
        return _merge_page_continuations(_pdf_elements(lines, levels, tables)), title


# --------------------------------------------------------------------------- files


def load_file(
    path: Path, *, filename: str | None = None, source_key: str | None = None
) -> Document:
    """Load a supported file into a :class:`Document`.

    ``filename`` overrides the display name (uploads arrive under temp names), and
    ``source_key`` sets the identity used for the ``doc_id`` (defaults to the filename).
    """
    filename = filename or path.name
    source_type = detect_source_type(filename)
    if not path.is_file():
        raise LoaderError(f"File not found: {filename}")

    meta_title: str | None = None
    if source_type == SourceType.PDF:
        elements, meta_title = load_pdf_elements(path)
    elif source_type == SourceType.DOCX:
        elements, meta_title = load_docx_elements(path)
    elif source_type == SourceType.MARKDOWN:
        elements, meta_title = parse_markdown(read_text_file(path))
    else:
        elements = parse_plain_text(read_text_file(path))

    if not any(e.kind != ElementKind.HEADING for e in elements):
        hint = (
            " Scanned PDFs need OCR, which isn't supported."
            if source_type == SourceType.PDF
            else ""
        )
        raise LoaderError(f"No extractable text found in {filename}.{hint}")

    key = normalize_source_key(source_key or filename)
    return Document(
        doc_id=make_doc_id(key),
        source=key,
        filename=filename,
        source_type=source_type,
        title=meta_title or infer_title(elements, title_from_filename(filename)),
        content_hash=file_sha256(path),
        elements=elements,
    )


# --------------------------------------------------------------------------- urls


def ensure_public_url(url: str) -> None:
    """Reject non-HTTP(S) URLs and hosts resolving to private/loopback/link-local addresses.

    This blocks server-side request forgery (e.g. the cloud metadata endpoint at
    169.254.169.254) from the "Add URL" feature. It checks every redirect hop. DNS
    rebinding between this check and the request is a residual risk, acceptable for
    this app's threat model.
    """
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https"):
        raise LoaderError("Only http:// and https:// URLs are supported.")
    host = parsed.hostname
    if not host:
        raise LoaderError(f"Invalid URL: {url}")
    port = parsed.port or (443 if parsed.scheme == "https" else 80)
    try:
        infos = socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
    except socket.gaierror as exc:
        raise LoaderError(f"Could not resolve host {host!r}.") from exc
    for info in infos:
        address = ipaddress.ip_address(str(info[4][0]).split("%", 1)[0])
        if not address.is_global:
            raise LoaderError(
                "URLs that point to private or local network addresses are not allowed."
            )


async def fetch_url(
    url: str,
    *,
    timeout_s: float = 20.0,
    max_bytes: int = 5 * 1024 * 1024,
    transport: httpx.AsyncBaseTransport | None = None,
) -> tuple[str, str]:
    """Download an HTML page with SSRF checks on every hop. Returns ``(html, final_url)``.

    ``transport`` lets tests substitute an in-memory HTTP transport.
    """
    current = url.strip()
    async with httpx.AsyncClient(
        follow_redirects=False,
        timeout=timeout_s,
        headers={"User-Agent": USER_AGENT},
        transport=transport,
    ) as client:
        for _ in range(MAX_REDIRECTS + 1):
            await asyncio.to_thread(ensure_public_url, current)
            try:
                async with client.stream("GET", current) as response:
                    if response.is_redirect:
                        current = urljoin(current, response.headers.get("location", ""))
                        continue
                    if response.status_code >= 400:
                        raise LoaderError(f"The page returned HTTP {response.status_code}.")
                    content_type = response.headers.get("content-type", "").lower()
                    if (
                        content_type
                        and "html" not in content_type
                        and "text/plain" not in content_type
                    ):
                        raise LoaderError(
                            f"The URL isn't an HTML page (content type {content_type})."
                        )
                    body = bytearray()
                    async for part in response.aiter_bytes():
                        body.extend(part)
                        if len(body) > max_bytes:
                            raise LoaderError("The page is too large to ingest.")
                    return body.decode(response.encoding or "utf-8", errors="replace"), current
            except httpx.HTTPError as exc:
                raise LoaderError(f"Could not download the page ({type(exc).__name__}).") from exc
    raise LoaderError("Too many redirects.")


def parse_html(html: str, url: str) -> tuple[list[DocumentElement], str | None, str]:
    """Extract main content with trafilatura. Returns ``(elements, title, markdown)``."""
    import trafilatura

    markdown = trafilatura.extract(
        html,
        url=url,
        output_format="markdown",
        include_tables=True,
        include_formatting=True,
        include_links=False,
        include_comments=False,
        favor_recall=True,
    )
    if not markdown or not markdown.strip():
        raise LoaderError("Couldn't find readable article content on that page.")
    elements, _ = parse_markdown(markdown)
    metadata = trafilatura.extract_metadata(html, default_url=url)
    title = (metadata.title or "").strip() if metadata is not None else ""
    return elements, title or None, markdown


def url_display_name(url: str) -> str:
    """Short, filename-like label for a URL: ``example.com/docs/setup``."""
    parsed = urlparse(url)
    return f"{parsed.hostname or ''}{parsed.path}".rstrip("/")[:120] or url[:120]


def document_from_html(html: str, url: str) -> Document:
    """Build a Document from downloaded HTML (pure; see :func:`fetch_url` for I/O)."""
    elements, title, markdown = parse_html(html, url)
    key = normalize_source_key(url)
    name = url_display_name(key)
    return Document(
        doc_id=make_doc_id(key),
        source=key,
        filename=name,
        source_type=SourceType.URL,
        title=title or infer_title(elements, name),
        # Hash the extracted text rather than raw HTML, so rotating ads or nonces in the
        # markup don't force a re-index.
        content_hash=sha256_text(markdown),
        elements=elements,
    )
