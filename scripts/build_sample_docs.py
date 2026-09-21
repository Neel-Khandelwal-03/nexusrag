"""Generate the binary sample documents in ``data/`` (PDF and DOCX).

The Markdown and TXT samples are checked in as-is. The PDFs and the DOCX are built
from the content below so they're reproducible and reviewable as text. All content
is fictional ("Skylark Dynamics" is not a real company).

Usage:
    python scripts/build_sample_docs.py
"""

from __future__ import annotations

import datetime as dt
import io
from pathlib import Path

import docx
import pymupdf
from docx.shared import Pt

DATA_DIR = Path(__file__).resolve().parents[1] / "data"
FIXED_PDF_DATE = "D:20260301000000Z"

CSS = """
body { font-family: serif; font-size: 11pt; line-height: 1.35; }
h1 { font-size: 22pt; margin-bottom: 6pt; }
h2 { font-size: 16pt; margin-top: 14pt; margin-bottom: 4pt; }
h3 { font-size: 13pt; margin-top: 10pt; margin-bottom: 3pt; }
p { margin-top: 3pt; margin-bottom: 5pt; }
table { border-collapse: separate; border-spacing: 0; margin-top: 4pt; margin-bottom: 8pt; }
th, td { border: 1px solid #333; padding: 3pt 6pt; font-size: 10pt; }
th { font-weight: bold; }
li { margin-bottom: 2pt; }
"""


def _table(rows: list[list[str]]) -> str:
    header, *body = rows
    head = "".join(f"<th>{cell}</th>" for cell in header)
    rest = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in row) + "</tr>" for row in body)
    return f"<table><tr>{head}</tr>{rest}</table>"


AURORA_HTML = f"""
<h1>Aurora X1 Field Drone Technical Specification</h1>
<p>Document SKD-SPEC-AX1, revision C, published March 2026 by Skylark Dynamics Engineering.
Skylark Dynamics is a fictional company; this document is sample data for NexusRAG.</p>

<h2>1 Overview</h2>
<p>The Aurora X1 is a rugged quadcopter designed for field inspection of power lines, pipelines
and wind turbines. It is built for crews that need to launch within two minutes of arriving on
site and keep flying in light rain and gusty wind. The X1 replaces the Aurora X0, which was
discontinued in 2024.</p>
<p>Target customers are utility companies and infrastructure inspection contractors. The X1 is
sold as a kit that includes the aircraft, two batteries, the Skylark GC-3 ground controller and a
hard-shell transport case.</p>

<h2>2 Hardware</h2>
<h3>2.1 Airframe</h3>
<p>The airframe is a carbon-fibre monocoque with folding arms. Unfolded, the diagonal wheelbase
is 780 mm; folded, the aircraft measures 410 x 290 x 180 mm. Take-off weight is 4.2 kg with one
battery and no payload, and the maximum take-off weight is 6.4 kg.</p>
<h3>2.2 Propulsion</h3>
<p>Four brushless outrunner motors drive 18-inch folding propellers. Each motor has its own
electronic speed controller with thermal protection. If a single motor fails above 15 metres,
the X1 can still complete a controlled landing on the remaining three motors in a degraded yaw
mode.</p>
<h3>2.3 Battery System</h3>
<p>The X1 uses hot-swappable smart batteries. Swapping a battery takes under 30 seconds and does
not require powering down the flight computer.</p>
{
    _table(
        [
            ["Property", "Value"],
            ["Chemistry", "Lithium-ion (21700 cells)"],
            ["Capacity", "12,000 mAh"],
            ["Nominal voltage", "44.4 V (12S)"],
            ["Energy", "533 Wh"],
            ["Charge time (0-100%)", "75 minutes with the CH-400 charger"],
            ["Rated cycles", "400 cycles to 80% capacity"],
            ["Battery weight", "1.9 kg"],
        ]
    )
}
<p>Batteries discharge themselves to 60% after 10 days without use to extend their service
life. For storage longer than three months, keep batteries between 15 and 25 degrees Celsius.</p>

<h2>3 Performance</h2>
{
    _table(
        [
            ["Metric", "Value"],
            ["Maximum flight time (no payload)", "46 minutes"],
            ["Flight time with 2 kg payload", "38 minutes"],
            ["Maximum range (line of sight)", "12 km"],
            ["Maximum horizontal speed", "72 km/h"],
            ["Maximum wind resistance", "12 m/s"],
            ["Service ceiling", "5,000 m above sea level"],
            ["Operating temperature", "-20 to 50 degrees Celsius"],
            ["Ingress protection", "IP55"],
        ]
    )
}
<p>Flight times are measured at sea level and 25 degrees Celsius in still air, landing with 10%
battery remaining. Expect flights to be roughly 15% shorter below 0 degrees Celsius.</p>

<h2>4 Sensors and Payload</h2>
<p>The standard payload is the Skylark VZ-20 gimbal camera with a 20-megapixel 1-inch sensor and
12x hybrid zoom. The optional TH-640 thermal module adds a 640 x 512 radiometric thermal sensor
for detecting hotspots on electrical equipment. The payload mount uses the Skylark QuickLock
interface and supports payloads of up to 2.2 kg.</p>
<p>Obstacle sensing uses six stereo vision pairs plus an upward-facing time-of-flight sensor,
giving omnidirectional detection from 0.5 m to 40 m. RTK positioning is available with the
optional RTK-2 module, which provides 1 cm + 1 ppm horizontal accuracy.</p>

<h2>5 Safety Features</h2>
<ul>
<li>Automatic return-to-home when the battery reaches the calculated reserve, or when the control
link is lost for more than 3 seconds.</li>
<li>Geofencing with a configurable altitude ceiling (default 120 m).</li>
<li>ADS-B receiver that warns the pilot about crewed aircraft within 10 km.</li>
<li>Emergency propeller stop, triggered by holding both control sticks inward for 2 seconds.</li>
<li>Parachute port compatible with the SkySafe P2 parachute system.</li>
</ul>

<h2>6 Maintenance</h2>
<p>Following the maintenance schedule is a condition of the warranty.</p>
{
    _table(
        [
            ["Interval", "Task"],
            [
                "Before every flight",
                "Inspect propellers for cracks; check battery latches; verify compass calibration",
            ],
            ["Every 50 flight hours", "Replace propellers; clean motor vents; update firmware"],
            [
                "Every 200 flight hours",
                "Replace motor bearings; inspect arm hinges; factory IMU calibration",
            ],
            ["Every 12 months", "Full service at an authorised Skylark service centre"],
        ]
    )
}

<h2>7 Warranty</h2>
<p>The Aurora X1 carries a 12-month limited warranty covering manufacturing defects. Batteries
are covered for 6 months or 150 cycles, whichever comes first. Crash damage is not covered, but
customers can buy Skylark Care Plus, which covers up to two accidental-damage replacements per
year for an annual fee of 890 USD.</p>
"""

QBR_HTML = f"""
<h1>Skylark Dynamics Q2 2026 Quarterly Business Review</h1>
<p>Prepared by the Finance and Strategy team for the board meeting of 21 July 2026. All figures
are in millions of US dollars unless stated otherwise. Skylark Dynamics is a fictional company;
this document is sample data for NexusRAG.</p>

<h2>1 Executive Summary</h2>
<p>Skylark Dynamics delivered record revenue of 48.6 million USD in Q2 2026, up 22% from
39.8 million USD in Q2 2025. Gross margin improved to 41.5% as the Borealis S2 moved to
higher-volume production. The motor supply constraint that slowed Aurora X1 deliveries in April
was resolved in May, and order fulfilment times are back within target.</p>
<p>Operating income doubled year over year to 4.8 million USD. The company opened its India
office in Bengaluru and launched the Borealis S2 LiDAR kit.</p>

<h2>2 Financial Highlights</h2>
{
    _table(
        [
            ["Metric", "Q2 2025", "Q2 2026", "Change"],
            ["Revenue", "39.8", "48.6", "+22%"],
            ["Gross margin", "38.9%", "41.5%", "+2.6 pts"],
            ["Operating expenses", "13.1", "15.4", "+18%"],
            ["Operating income", "2.4", "4.8", "+100%"],
            ["Headcount (end of quarter)", "412", "486", "+18%"],
        ]
    )
}

<h2>3 Revenue by Product Line</h2>
{
    _table(
        [
            ["Product line", "Units shipped", "Revenue", "Share of revenue"],
            ["Aurora X1", "1,420", "20.6", "42%"],
            ["Borealis S2", "610", "11.5", "24%"],
            ["Payloads and accessories", "n/a", "7.9", "16%"],
            ["Software and services (Skylark Cloud, Care Plus)", "n/a", "8.6", "18%"],
        ]
    )
}
<p>The Aurora X1 lists at 14,500 USD per kit and remains the largest product line. The Borealis
S2 standard kit lists at 18,900 USD and the LiDAR kit at 31,500 USD. Software and services grew
fastest, up 41% year over year, driven by Skylark Cloud subscriptions from utility customers.</p>

<h2>4 Regional Performance</h2>
<p>North America generated 58% of revenue, Europe 27% and Asia-Pacific 15%. Europe grew fastest
at 31% year over year after the Borealis S2 received its EU class C2 certification in March. The
new Bengaluru office opened in April 2026 with 35 staff focused on flight software and customer
support for the Asia-Pacific region.</p>

<h2>5 Product Updates</h2>
<h3>5.1 Aurora X1</h3>
<p>Firmware 4.2 shipped in May with improved ADS-B alerting and a new automated wind-turbine
inspection mode that plans the flight path around the blades. The TH-640 thermal module became
the most popular accessory, attached to 38% of Aurora X1 kits sold in the quarter.</p>
<h3>5.2 Borealis S2</h3>
<p>The Borealis S2 LiDAR kit launched in June 2026. Early customers in mining report surveying
twice the area per day compared with their previous multicopter LiDAR systems.</p>
<h3>5.3 Roadmap</h3>
<p>The Aurora X2 is in development with a planned launch in Q2 2027. Its specifications have not
been disclosed.</p>

<h2>6 Operations</h2>
<p>In April, motor supplier Voltra Motion delayed shipments by five weeks, which reduced Aurora
X1 output. The team qualified a second motor supplier in May. Average order fulfilment time fell
from 9 days at the end of April to 4 days at the end of June. Customer satisfaction (CSAT) was
4.5 out of 5, while support ticket volume rose 12% with the larger installed base.</p>

<h2>7 Outlook for Q3 2026</h2>
<p>Revenue guidance for Q3 2026 is 51 to 54 million USD. The company plans to hire 40 engineers,
mainly for the Aurora X2 programme, and expects gross margin to stay between 40% and 42%.</p>

<h2>8 Key Risks</h2>
<ul>
<li>Tariffs on imported battery cells could raise battery costs by up to 12%.</li>
<li>Changes to FAA rules on beyond-visual-line-of-sight (BVLOS) operations could delay enterprise
deployments in the United States.</li>
<li>Dependence on a single supplier for the VZ-20 camera sensor.</li>
</ul>
"""


def build_pdf(path: Path, html: str, title: str, footer: str) -> None:
    """Flow HTML across A4 pages, then stamp a running footer and set metadata."""
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (56, 56, -56, -64)  # noqa: RUF005 (Rect arithmetic)
    buffer = io.BytesIO()
    story = pymupdf.Story(html=html, user_css=CSS)
    writer = pymupdf.DocumentWriter(buffer)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()

    with pymupdf.open("pdf", buffer.getvalue()) as doc:
        for number, page in enumerate(doc, start=1):
            page.insert_text(
                (56, mediabox.height - 32),
                f"{footer} | Page {number} of {doc.page_count}",
                fontsize=8,
                color=(0.4, 0.4, 0.4),
            )
        doc.set_metadata(
            {
                "title": title,
                "author": "Skylark Dynamics (fictional sample data)",
                "creationDate": FIXED_PDF_DATE,
                "modDate": FIXED_PDF_DATE,
            }
        )
        doc.save(str(path), garbage=3, deflate=True)


def build_borealis_docx(path: Path) -> None:
    document = docx.Document()
    document.styles["Normal"].font.size = Pt(11)
    props = document.core_properties
    props.title = "Borealis S2 Survey Drone Product Sheet"
    props.author = "Skylark Dynamics (fictional sample data)"
    props.created = props.modified = dt.datetime(2026, 3, 1, tzinfo=dt.UTC)

    def table(rows: list[list[str]]) -> None:
        t = document.add_table(rows=0, cols=len(rows[0]))
        t.style = "Table Grid"
        for row in rows:
            cells = t.add_row().cells
            for cell, value in zip(cells, row, strict=True):
                cell.text = value

    add, para = document.add_heading, document.add_paragraph
    add("Borealis S2 Survey Drone Product Sheet", 0)
    para(
        "Product sheet SKD-PS-BS2, published March 2026 by Skylark Dynamics. Skylark Dynamics is "
        "a fictional company; this document is sample data for NexusRAG."
    )

    add("1 Overview", 1)
    para(
        "The Borealis S2 is a fixed-wing VTOL survey drone for large-area mapping, agriculture and "
        "mining surveys. It takes off vertically like a multicopter, then transitions to efficient "
        "fixed-wing flight. The S2 launched in September 2025 and is designed for surveyors who map "
        "up to 1,000 hectares in a single flight."
    )

    add("2 Hardware", 1)
    add("2.1 Airframe", 2)
    para(
        "The S2 has a 2.1 m wingspan and a length of 1.2 m, built from EPO foam reinforced with a "
        "carbon composite spar. Take-off weight is 7.8 kg and maximum take-off weight is 9.5 kg. "
        "The wings attach without tools, and the aircraft can be assembled in under 5 minutes."
    )
    add("2.2 Battery System", 2)
    table(
        [
            ["Property", "Value"],
            ["Chemistry", "Semi-solid-state lithium-ion"],
            ["Capacity", "22,000 mAh"],
            ["Nominal voltage", "22.2 V (6S)"],
            ["Energy", "488 Wh"],
            ["Charge time (0-100%)", "110 minutes with the CH-600 charger"],
            ["Rated cycles", "600 cycles to 80% capacity"],
            ["Battery weight", "2.6 kg"],
        ]
    )
    para(
        "Unlike the Aurora X1, the Borealis S2 battery is not hot-swappable: the aircraft must be "
        "powered off before the battery is changed."
    )

    add("3 Performance", 1)
    table(
        [
            ["Metric", "Value"],
            ["Maximum flight time", "95 minutes"],
            ["Maximum range (line of sight)", "18 km"],
            ["Maximum range (with LTE relay)", "60 km"],
            ["Cruise speed", "64 km/h"],
            ["Maximum horizontal speed", "90 km/h"],
            ["Maximum wind resistance", "10 m/s"],
            ["Service ceiling", "4,500 m above sea level"],
            ["Operating temperature", "-10 to 45 degrees Celsius"],
            ["Ingress protection", "IP43"],
            ["Coverage per flight at 3 cm/px", "1,000 hectares"],
        ]
    )

    add("4 Sensors and Payload", 1)
    para(
        "The standard payload is the MP-45 mapping camera, a 45-megapixel full-frame camera with a "
        "mechanical shutter. The optional LX-1 LiDAR module captures 240,000 points per second with "
        "1.5 cm vertical accuracy. RTK and PPK GNSS are built in, so no extra positioning module is "
        "needed. Payload capacity is 1.2 kg."
    )

    add("5 Safety Features", 1)
    for item in (
        "Automatic return-to-home when the control link is lost for more than 5 seconds.",
        "Geofencing with a configurable altitude ceiling (default 120 m).",
        "ADS-B receiver that warns the pilot about crewed aircraft.",
        "Automatic transition abort: if the switch to fixed-wing flight fails, the aircraft "
        "returns to hover mode.",
        "The S2 has no parachute port.",
    ):
        para(item, style="List Bullet")

    add("6 Maintenance", 1)
    table(
        [
            ["Interval", "Task"],
            [
                "Before every flight",
                "Check control surfaces; inspect VTOL propellers; confirm the airspeed sensor is clear",
            ],
            [
                "Every 100 flight hours",
                "Replace VTOL propellers; update firmware; check servo linkages",
            ],
            ["Every 300 flight hours", "Replace the cruise motor; factory calibration"],
            ["Every 12 months", "Full service at an authorised Skylark service centre"],
        ]
    )

    add("7 Warranty", 1)
    para(
        "The Borealis S2 carries a 24-month limited warranty covering manufacturing defects. "
        "Batteries are covered for 12 months or 300 cycles, whichever comes first. Skylark Care Plus "
        "is available for 1,250 USD per year and covers one accidental-damage replacement per year."
    )

    add("8 Pricing", 1)
    para(
        "The standard Borealis S2 kit costs 18,900 USD. The LiDAR kit, which adds the LX-1 module "
        "and a second battery, costs 31,500 USD."
    )

    document.save(str(path))


def main() -> None:
    DATA_DIR.mkdir(exist_ok=True)
    build_pdf(
        DATA_DIR / "aurora-x1-spec.pdf",
        AURORA_HTML,
        "Aurora X1 Field Drone Technical Specification",
        "Skylark Dynamics | SKD-SPEC-AX1 rev C",
    )
    build_pdf(
        DATA_DIR / "q2-2026-business-review.pdf",
        QBR_HTML,
        "Skylark Dynamics Q2 2026 Quarterly Business Review",
        "Skylark Dynamics | Internal",
    )
    build_borealis_docx(DATA_DIR / "borealis-s2-product-sheet.docx")
    for path in sorted(DATA_DIR.iterdir()):
        print(f"{path.name:<36} {path.stat().st_size:>8,} bytes")


if __name__ == "__main__":
    main()
