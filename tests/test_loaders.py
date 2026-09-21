from __future__ import annotations

import socket
from pathlib import Path

import docx
import httpx
import pymupdf
import pytest

from nexusrag.ingestion.loaders import (
    LoaderError,
    UnsupportedFileError,
    document_from_html,
    ensure_public_url,
    fetch_url,
    load_file,
    parse_markdown,
    parse_plain_text,
    table_to_markdown,
)
from nexusrag.models import ElementKind, SourceType

# --------------------------------------------------------------------------- markdown


def kinds(elements: list) -> list[tuple[str, int | None]]:
    return [(e.kind.value, e.level) for e in elements]


def test_markdown_structure() -> None:
    text = """---
title: Front Matter Title
---

# Main

Intro paragraph
that wraps.

Setext Heading
--------------

```python
# not a heading
x = 1
```

| A | B |
|---|---|
| 1 | 2 |

- item one
  continued
- item two

***

> quoted text
"""
    elements, title = parse_markdown(text)
    assert title == "Front Matter Title"
    assert kinds(elements) == [
        ("heading", 1),
        ("paragraph", None),
        ("heading", 2),
        ("code", None),
        ("table", None),
        ("list", None),
        ("paragraph", None),
    ]
    assert elements[1].text == "Intro paragraph that wraps."
    assert "# not a heading" in elements[3].text
    assert elements[4].text.splitlines()[2] == "| 1 | 2 |"
    assert elements[5].text == "- item one\n  continued\n- item two"
    assert elements[6].text == "quoted text"


def test_plain_text_headings_and_lists() -> None:
    text = "USER GUIDE\n\n1.2 Charging the battery\n\nPlug it in\nand wait.\n\n- one\n- two"
    elements = parse_plain_text(text)
    assert kinds(elements) == [("heading", 1), ("heading", 2), ("paragraph", None), ("list", None)]
    assert elements[2].text == "Plug it in and wait."


def test_plain_text_sentence_is_not_a_heading() -> None:
    assert kinds(parse_plain_text("This short line ends with a period.")) == [("paragraph", None)]


def test_table_to_markdown() -> None:
    md = table_to_markdown([["Name", "Note"], ["a|b", None], ["c", "multi\nline", "extra"]])
    assert md.splitlines() == [
        "| Name | Note |  |",
        "|---|---|---|",
        "| a\\|b |  |  |",
        "| c | multi line | extra |",
    ]
    assert table_to_markdown([["only one column"], ["x"]]) == ""
    assert table_to_markdown([["header", "only"]]) == ""
    assert table_to_markdown([["header", "only"]], min_rows=1).count("\n") == 1


# --------------------------------------------------------------------------- files


def test_markdown_file(tmp_path: Path) -> None:
    path = tmp_path / "guide.md"
    path.write_text("# Guide\n\n## Setup\n\nInstall it.\n", encoding="utf-8")
    doc = load_file(path)
    assert doc.title == "Guide"
    assert doc.source_type == SourceType.MARKDOWN
    assert doc.source == "guide.md"
    assert len(doc.content_hash) == 64


def test_source_key_controls_doc_id(tmp_path: Path) -> None:
    path = tmp_path / "upload_123.tmp.md"
    path.write_text("Hello there.", encoding="utf-8")
    a = load_file(path, filename="notes.md")
    b = load_file(path, filename="notes.md", source_key="team/notes.md")
    assert a.filename == "notes.md"
    assert a.doc_id != b.doc_id
    assert a.title == "Notes"


def test_unsupported_and_empty_files(tmp_path: Path) -> None:
    exe = tmp_path / "tool.exe"
    exe.write_bytes(b"MZ")
    with pytest.raises(UnsupportedFileError):
        load_file(exe)
    empty = tmp_path / "empty.txt"
    empty.write_text("   \n", encoding="utf-8")
    with pytest.raises(LoaderError, match="No extractable text"):
        load_file(empty)


def test_cp1252_text_fallback(tmp_path: Path) -> None:
    path = tmp_path / "legacy.txt"
    path.write_bytes("Caf\xe9 menu costs 5\x80.".encode("latin-1"))
    assert "Café" in load_file(path).elements[0].text


def test_docx(tmp_path: Path) -> None:
    document = docx.Document()
    document.core_properties.title = "Product Sheet"
    document.add_heading("Product Sheet", 0)
    document.add_heading("1 Overview", 1)
    document.add_paragraph("A survey drone.")
    document.add_paragraph("First point", style="List Bullet")
    document.add_paragraph("Second point", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Metric", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "Range", "18 km"
    path = tmp_path / "sheet.docx"
    document.save(str(path))

    doc = load_file(path)
    assert doc.title == "Product Sheet"
    assert kinds(doc.elements) == [
        ("heading", 1),
        ("paragraph", None),
        ("list", None),
        ("table", None),
    ]
    assert doc.elements[2].text == "- First point\n- Second point"
    assert "| Range | 18 km |" in doc.elements[3].text


def _story_pdf(path: Path, html: str, footer: str | None = None) -> None:
    """Render HTML to a multi-page PDF with an optional running footer."""
    import io

    buffer = io.BytesIO()
    mediabox = pymupdf.paper_rect("a4")
    story = pymupdf.Story(
        html=html,
        user_css="h1{font-size:22pt} h2{font-size:16pt} p{font-size:11pt} "
        "table{border-collapse:separate;border-spacing:0} "
        "td,th{border:1px solid black;padding:2pt}",
    )
    writer = pymupdf.DocumentWriter(buffer)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(mediabox + (56, 56, -56, -64))  # noqa: RUF005 (Rect arithmetic)
        story.draw(device)
        writer.end_page()
    writer.close()
    with pymupdf.open("pdf", buffer.getvalue()) as doc:
        if footer:
            for number, page in enumerate(doc, start=1):
                page.insert_text(
                    (56, mediabox.height - 30), f"{footer} | Page {number}", fontsize=8
                )
        doc.set_metadata({"title": "Meta Title"})
        doc.save(str(path))


def test_pdf_structure_tables_and_furniture(tmp_path: Path) -> None:
    filler = "".join(
        f"<p>Paragraph {i} about field inspections and batteries.</p>" for i in range(45)
    )
    html = (
        "<h1>Drone Manual</h1><h2>1 Overview</h2><p>The drone flies far.</p>"
        "<table><tr><th>Metric</th><th>Value</th></tr><tr><td>Range</td><td>12 km</td></tr></table>"
        f"<h2>2 Details</h2>{filler}<h2>3 Warranty</h2><p>Twelve months of cover.</p>"
    )
    path = tmp_path / "manual.pdf"
    _story_pdf(path, html, footer="ACME Confidential")
    doc = load_file(path)

    assert doc.title == "Meta Title"
    headings = [(e.text, e.level) for e in doc.elements if e.kind == ElementKind.HEADING]
    assert headings == [("Drone Manual", 1), ("1 Overview", 2), ("2 Details", 2), ("3 Warranty", 2)]
    tables = [e for e in doc.elements if e.kind == ElementKind.TABLE]
    assert len(tables) == 1
    assert "| Range | 12 km |" in tables[0].text
    assert tables[0].page == 1
    text = "\n".join(e.text for e in doc.elements)
    assert "ACME Confidential" not in text  # running footer removed
    assert "Range 12 km" not in text  # table text not duplicated as a paragraph
    assert max(e.page or 0 for e in doc.elements) >= 2


def test_pdf_without_text_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    with pymupdf.open() as doc:
        doc.new_page()
        doc.save(str(path))
    with pytest.raises(LoaderError, match="OCR"):
        load_file(path)


def test_corrupt_pdf(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a pdf at all")
    with pytest.raises(LoaderError):
        load_file(path)


# --------------------------------------------------------------------------- urls


@pytest.fixture
def fake_dns(monkeypatch: pytest.MonkeyPatch) -> dict[str, str]:
    table = {
        "example.com": "93.184.216.34",
        "internal.corp": "10.0.0.5",
        "meta.local": "169.254.169.254",
    }

    def getaddrinfo(host: str, port: int, *args: object, **kwargs: object) -> list:
        if host not in table and host != "127.0.0.1":
            raise socket.gaierror("unknown host")
        address = table.get(host, host)
        return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (address, port))]

    monkeypatch.setattr(socket, "getaddrinfo", getaddrinfo)
    return table


@pytest.mark.usefixtures("fake_dns")
@pytest.mark.parametrize(
    "url",
    [
        "ftp://example.com/file",
        "http://internal.corp/admin",
        "http://meta.local/latest/meta-data",
        "http://127.0.0.1:8080/",
        "https://nope.invalid/",
    ],
)
def test_ssrf_guard_rejects(url: str) -> None:
    with pytest.raises(LoaderError):
        ensure_public_url(url)


@pytest.mark.usefixtures("fake_dns")
def test_ssrf_guard_allows_public_host() -> None:
    ensure_public_url("https://example.com/docs")


ARTICLE = (
    "<html><head><title>Setup Guide</title></head><body><nav>menu</nav><article>"
    "<h1>Setup Guide</h1><p>" + "Install the app carefully before first use. " * 12 + "</p>"
    "<h2>Pairing</h2><p>" + "Hold the button for five seconds to pair. " * 12 + "</p>"
    "</article></body></html>"
)


@pytest.mark.usefixtures("fake_dns")
async def test_fetch_url_follows_public_redirects() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/old":
            return httpx.Response(301, headers={"location": "/new"})
        return httpx.Response(200, text=ARTICLE, headers={"content-type": "text/html"})

    html, final = await fetch_url("https://example.com/old", transport=httpx.MockTransport(handler))
    assert final == "https://example.com/new"
    assert "Setup Guide" in html


@pytest.mark.usefixtures("fake_dns")
async def test_fetch_url_blocks_redirect_to_private_address() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(302, headers={"location": "http://internal.corp/secrets"})

    with pytest.raises(LoaderError, match="private"):
        await fetch_url("https://example.com/", transport=httpx.MockTransport(handler))


@pytest.mark.usefixtures("fake_dns")
async def test_fetch_url_limits() -> None:
    big = httpx.MockTransport(
        lambda r: httpx.Response(200, text="x" * 5000, headers={"content-type": "text/html"})
    )
    with pytest.raises(LoaderError, match="too large"):
        await fetch_url("https://example.com/", max_bytes=1000, transport=big)
    pdf = httpx.MockTransport(
        lambda r: httpx.Response(200, content=b"%PDF", headers={"content-type": "application/pdf"})
    )
    with pytest.raises(LoaderError, match="isn't an HTML page"):
        await fetch_url("https://example.com/", transport=pdf)
    missing = httpx.MockTransport(lambda r: httpx.Response(404))
    with pytest.raises(LoaderError, match="404"):
        await fetch_url("https://example.com/", transport=missing)


def test_document_from_html() -> None:
    doc = document_from_html(ARTICLE, "https://example.com/guide#section")
    assert doc.source_type == SourceType.URL
    assert doc.source == "https://example.com/guide"
    assert doc.filename == "example.com/guide"
    assert doc.title == "Setup Guide"
    assert any(e.kind == ElementKind.HEADING and e.text == "Pairing" for e in doc.elements)
    assert "menu" not in doc.text  # boilerplate navigation removed
    assert document_from_html(ARTICLE, "https://example.com/guide").content_hash == doc.content_hash


def test_document_from_html_without_content() -> None:
    with pytest.raises(LoaderError):
        document_from_html("<html><body></body></html>", "https://example.com/")
